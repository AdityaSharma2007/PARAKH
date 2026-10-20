import pytest
from pathlib import Path
import fitz
import docx
from backend.app.services.resume_parser import ResumeParser
from backend.app.core.exceptions import ResumeParsingError

def test_extract_text_from_pdf(tmp_path):
    # Create sample PDF using PyMuPDF
    pdf_path = tmp_path / "sample_resume.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "John Doe\nSoftware Engineer\nSkills: Python, FastAPI, SQL\nExperience: Built microservices.")
    doc.save(str(pdf_path))
    doc.close()

    text = ResumeParser.parse_file(pdf_path)
    assert "John Doe" in text
    assert "FastAPI" in text

def test_extract_text_from_docx(tmp_path):
    # Create sample DOCX
    docx_path = tmp_path / "sample_resume.docx"
    doc = docx.Document()
    doc.add_heading("Jane Developer", 0)
    doc.add_paragraph("Skills: JavaScript, TypeScript, React, Docker")
    doc.save(str(docx_path))

    text = ResumeParser.parse_file(docx_path)
    assert "Jane Developer" in text
    assert "React" in text

def test_unsupported_file_extension(tmp_path):
    txt_path = tmp_path / "resume.txt"
    txt_path.write_text("Hello world")
    with pytest.raises(ResumeParsingError):
        ResumeParser.parse_file(txt_path)
