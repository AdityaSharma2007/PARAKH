try:
    import pymupdf as fitz  # PyMuPDF
except ImportError:
    import fitz
import docx
from pathlib import Path
from backend.app.core.exceptions import ResumeParsingError

class ResumeParser:
    @staticmethod
    def sanitize_text(text: str) -> str:
        if not text:
            return ""
        replacements = {
            "\u2013": "-",
            "\u2014": "-",
            "\u2015": "-",
            "\u2018": "'",
            "\u2019": "'",
            "\u201c": '"',
            "\u201d": '"',
            "\u2022": "*",
            "\u00a0": " ",
            "\ufffd": "-"
        }
        for k, v in replacements.items():
            text = text.replace(k, v)
        return text

    @staticmethod
    def extract_text_from_pdf(file_path: Path | str) -> str:
        try:
            doc = fitz.open(str(file_path))
            if doc.is_encrypted:
                raise ResumeParsingError("PDF file is encrypted or password-protected.")
            
            text_chunks = []
            for page in doc:
                text_chunks.append(page.get_text("text"))
            
            full_text = "\n".join(text_chunks).strip()
            if not full_text:
                raise ResumeParsingError(
                    "No readable text could be extracted from this PDF. It may be scanned or image-only."
                )
            return ResumeParser.sanitize_text(full_text)
        except fitz.FileDataError as e:
            raise ResumeParsingError(f"Corrupted or invalid PDF file: {e}")
        except Exception as e:
            if isinstance(e, ResumeParsingError):
                raise
            raise ResumeParsingError(f"Failed to extract PDF contents: {e}")

    @staticmethod
    def extract_text_from_docx(file_path: Path | str) -> str:
        try:
            doc = docx.Document(str(file_path))
            text_chunks = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    text_chunks.append(para.text.strip())
            
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_chunks.append(row_text)
            
            full_text = "\n".join(text_chunks).strip()
            if not full_text:
                raise ResumeParsingError("No readable text found in DOCX file.")
            return ResumeParser.sanitize_text(full_text)
        except Exception as e:
            raise ResumeParsingError(f"Failed to extract DOCX contents: {e}")

    @classmethod
    def parse_file(cls, file_path: Path | str) -> str:
        path = Path(file_path)
        ext = path.suffix.lower()
        if ext == ".pdf":
            return cls.extract_text_from_pdf(path)
        elif ext == ".docx":
            return cls.extract_text_from_docx(path)
        else:
            raise ResumeParsingError(f"Unsupported file format: {ext}")

    @staticmethod
    def extract_contact_info(text: str) -> dict:
        import re
        contacts = {
            "email": None,
            "phone": None,
            "linkedin": None,
            "github": None,
            "portfolio": None
        }

        # Email
        email_match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
        if email_match:
            contacts["email"] = email_match.group(0).strip().rstrip(".,;:")

        # Phone
        phone_match = re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text)
        if phone_match:
            contacts["phone"] = phone_match.group(0).strip().rstrip(".,;:")

        # LinkedIn
        li_match = re.search(r"(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/([a-zA-Z0-9_-]+)", text, re.IGNORECASE)
        if li_match:
            contacts["linkedin"] = f"https://linkedin.com/in/{li_match.group(1)}".rstrip(".,;:")

        # GitHub
        gh_match = re.search(r"(?:https?:\/\/)?(?:www\.)?github\.com\/([a-zA-Z0-9_-]+)", text, re.IGNORECASE)
        if gh_match:
            contacts["github"] = f"https://github.com/{gh_match.group(1)}".rstrip(".,;:")

        # Portfolio / Website
        web_match = re.search(r"(?:https?:\/\/)(?!.*(?:linkedin|github|google))[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}(?:\/[^\s]*)?", text, re.IGNORECASE)
        if web_match:
            contacts["portfolio"] = web_match.group(0).strip().rstrip(".,;:")

        return contacts

    @staticmethod
    def segment_sections(text: str) -> dict:
        import re
        section_headers = {
            "SUMMARY": r"(?:summary|professional summary|executive summary|objective|about me|profile)",
            "SKILLS": r"(?:technical skills|skills|technologies|core competencies|competencies|areas of expertise|tech stack)",
            "PROJECTS": r"(?:projects|personal projects|academic projects|key projects|notable projects|recent projects)",
            "EXPERIENCE": r"(?:experience|work experience|professional experience|employment history|work history|internships)",
            "EDUCATION": r"(?:education|academic background|academic qualifications|academics)",
            "CERTIFICATIONS": r"(?:certifications|licenses & certifications|certificates|awards & achievements|achievements)"
        }

        lines = text.splitlines()
        found_sections = {}
        current_section = "HEADER"
        current_lines = []

        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                current_lines.append("")
                continue

            # Check if this line looks like a header (short, matches section keywords)
            matched_header = None
            if len(trimmed) < 45 and not trimmed.endswith("."):
                clean_line = re.sub(r"[^a-zA-Z\s]", "", trimmed).strip().lower()
                for sec_name, pattern in section_headers.items():
                    if re.fullmatch(pattern, clean_line):
                        matched_header = sec_name
                        break

            if matched_header:
                if current_lines:
                    found_sections[current_section] = "\n".join(current_lines).strip()
                current_section = matched_header
                current_lines = []
            else:
                current_lines.append(line)

        if current_lines:
            found_sections[current_section] = "\n".join(current_lines).strip()

        return found_sections

